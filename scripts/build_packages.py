"""Build the Pantons operator kit and its update ZIP. Maintainer utility only."""
from pathlib import Path
from datetime import datetime
import argparse
import hashlib
import json
import platform
import re
import shutil
import zipfile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]
VERSION = (ROOT / 'packaging/version.txt').read_text().strip()
PERSISTENT = json.loads((ROOT / 'packaging/persistent-artifacts.json').read_text())
REPRO = json.loads((ROOT / 'packaging/reproducibility.json').read_text())
WEEKLY_DIRS = PERSISTENT['artifacts']['weekly_folder_layout']['directories']
METADATA_SENTINEL = '1980-01-01T00:00:00Z'
METADATA_TIMESTAMP = datetime.fromisoformat(REPRO.get('metadata_timestamp', '').replace('Z','+00:00')) if REPRO.get('metadata_timestamp') else None
ZIP_TIMESTAMP = ((METADATA_TIMESTAMP.year, METADATA_TIMESTAMP.month, METADATA_TIMESTAMP.day,
                  METADATA_TIMESTAMP.hour, METADATA_TIMESTAMP.minute, METADATA_TIMESTAMP.second)
                 if METADATA_TIMESTAMP else None)

# Operator kit layout.
KIT = 'Panton Scheduling Assistant'
FULL_ZIP = 'Panton_Scheduling_Assistant.zip'
UPDATE_ZIP = 'Update.zip'
ZIPS = (FULL_ZIP, UPDATE_ZIP)
TASK_DIR = 'TASK FILES'
WEEK_TEMPLATE = 'WEEKS/_COPY ME - New week'
START_DOC = 'START HERE.docx'
RULES_DOC = 'Squadron Rules.docx'
INPUT_CHECKLIST = 'INPUT CHECKLIST.docx'
MIGRATION_DOC = 'READ ME - Squadron Rules changes.docx'
PROFILE_STYLE_DOCS = (RULES_DOC,)

# Maintained sources.
KIT_SOURCES = ROOT / 'packaging/kit'
CORE_SOURCE = KIT_SOURCES / 'core.md'
WEEK_FILE_SOURCE = KIT_SOURCES / 'week-file.md'
START_SOURCE = KIT_SOURCES / 'start-here.md'
INPUT_SOURCE = KIT_SOURCES / 'input-checklist.md'
RULES_SOURCE = KIT_SOURCES / 'squadron-rules.md'
PROFILE_SOURCE = ROOT / 'local-profiles/pantons.md'
TASKS = [
    ('1-tuesday-plan.md', '1 TUESDAY - Plan next week'),
    ('2-wednesday-check-draft.md', '2 WEDNESDAY - Check my draft'),
    ('3-thursday-full-qc.md', '3 THURSDAY - Full QC'),
    ('4-thursday-buy-sell.md', '4 THURSDAY - Buy-Sell'),
    ('5-thu-fri-changes-publish.md', '5 THU-FRI - Check changes and publish'),
    ('6-execution-week-change.md', '6 EXECUTION WEEK - Schedule change'),
    ('do-only-independent-review.md', 'DO ONLY - Independent review'),
    ('help-next.md', 'HELP - What do I do next'),
]

# Task-file contract.
DELEGATION_GUARD = 'Do not delegate to subagents.'
GUARD = ('Do not delegate to subagents. Do all of this work yourself in this chat: do not hand any part of it to '
         'another model or agent, launch parallel assistants, or use automatic agent delegation. Looking up '
         'publications in DoW Policies Beta (GAMECHANGER) is allowed; it is evidence retrieval, not delegation. '
         'No specific model is required.')
END_MARKER = '=== END OF TASK FILE ==='
TASK_FILE_LIMIT = 40000
NAMED_MODEL = re.compile(r'\bGemini\b|\bTerra\b|\bGrok\b|GPT-?\d', re.I)
LINK_NAMES = {}


def task_file_text(source, title):
    """Assemble one self-contained task file: guard, header, core, task script, WEEK FILE format, end marker."""
    header = ('=== PANTON SCHEDULING ASSISTANT - TASK FILE ===\n'
              f'THIS TASK: {title}\n'
              f'Kit version: {VERSION}\n'
              'Read this whole file before you reply. It has three parts: Part 1, core instructions that are the same '
              "in every task file; Part 2, this task's step-by-step script; Part 3, the WEEK FILE format. It ends "
              'with an end marker.')
    parts = [GUARD, header, CORE_SOURCE.read_text().strip(), (KIT_SOURCES / 'tasks' / source).read_text().strip(),
             WEEK_FILE_SOURCE.read_text().strip(), END_MARKER]
    return '\n\n'.join(parts) + '\n'


def validate_task_file(text, name):
    if not text.startswith(DELEGATION_GUARD):
        raise ValueError('Task file must start with the no-delegation instruction: ' + name)
    lines = text.rstrip('\n').splitlines()
    if not lines or lines[-1] != END_MARKER:
        raise ValueError('Task file must end with the END marker: ' + name)
    if len(text.encode('utf-8')) > TASK_FILE_LIMIT:
        raise ValueError(f'Task file exceeds the {TASK_FILE_LIMIT}-byte size budget: ' + name)
    if NAMED_MODEL.search(text):
        raise ValueError('Task file must require no named model: ' + name)
    if '{{' in text or '}}' in text:
        raise ValueError('Task file has an unfilled template field: ' + name)
    try:
        text.encode('ascii')
    except UnicodeEncodeError as exc:
        raise ValueError('Task file must be plain ASCII: ' + name) from exc


def squadron_rules_source():
    """Squadron Rules = the maintained frame with the approved Pantons profile as Section 1."""
    profile = PROFILE_SOURCE.read_text()
    body = profile.split('\n', 1)[1].strip()
    body = re.sub(r'^## ', '### ', body, flags=re.M)
    frame = RULES_SOURCE.read_text()
    if frame.count('{{PROFILE}}') != 1:
        raise ValueError('packaging/kit/squadron-rules.md must contain exactly one {{PROFILE}} marker')
    return frame.replace('{{PROFILE}}', body)


def validate_reproducibility_contract():
    if REPRO.get('schema_version') != 1 or REPRO.get('package_format') != 1:
        raise ValueError('Unsupported packaging/reproducibility.json contract')
    running=platform.python_version()
    if running != REPRO.get('python_version'):
        raise ValueError(f'Package build requires Python {REPRO.get("python_version")}; running {running}')
    if REPRO.get('archive_mode') != 'stored':
        raise ValueError('Current package format requires archive_mode=stored')
    if REPRO.get('metadata_timestamp') != METADATA_SENTINEL or ZIP_TIMESTAMP != (1980,1,1,0,0,0):
        raise ValueError('Package format 1 requires the documented 1980-01-01 metadata sentinel')
    requirements=(ROOT/'packaging/requirements.txt').read_bytes()
    if hashlib.sha256(requirements).hexdigest() != REPRO.get('requirements_sha256'):
        raise ValueError('packaging/requirements.txt changed without updating reproducibility.json')


def actionable_migrations():
    superseded={prior for item in PERSISTENT['migrations'] for prior in item.get('supersedes', [])}
    return [item for item in PERSISTENT['migrations'] if item['action'] != 'none' and item['id'] not in superseded]


def _demote(source):
    return re.sub(r'^(#{1,5}) ', lambda m: '#' + m.group(1) + ' ', source, flags=re.M)


def migration_reference_sources():
    """Reference-only copies for active migrations, keyed by title. Squadron Rules is shown fully assembled."""
    refs = {}
    for item in actionable_migrations():
        for artifact in item['artifacts']:
            spec = PERSISTENT['artifacts'][artifact]
            if 'source' not in spec:
                continue
            if artifact in ('pantons_local_profile', 'squadron_rules_frame'):
                refs.setdefault('Squadron Rules', squadron_rules_source())
            elif artifact == 'weekly_input_checklist':
                refs.setdefault('Input checklist', INPUT_SOURCE.read_text())
            else:
                raise ValueError('No reference copy rule for persistent artifact: ' + artifact)
    return refs


def persistent_migration_source():
    migrations = actionable_migrations()
    if not migrations:
        return ''
    lines = [
        '# Squadron Rules changes',
        '',
        'This file comes with an update only when a release changes Squadron Rules, the input checklist or the weekly folder. '
        'Nothing in your Squadron Rules or WEEKS was changed automatically. Follow each item that applies to you, '
        'have the DO confirm any change to Squadron Rules, then delete this file.',
        '',
    ]
    for item in migrations:
        lines += [f'## {item["id"]} (package {item["introduced_in"]})', '', item['summary'], '']
        lines += ['- [ ] ' + instruction for instruction in item['instructions']]
        lines += ['']
    for title, source in migration_reference_sources().items():
        lines += [f'## Reference copy: {title}', '',
                  'Comparison material only. It is not an approved replacement for your current file.', '',
                  _demote(_demote(source)).strip(), '']
    return '\n'.join(lines)


def el(tag, **attrs):
    node = OxmlElement('w:'+tag)
    for k,v in attrs.items(): node.set(qn('w:'+k), str(v))
    return node


def configure(doc, checklist=False):
    """compact_reference_guide preset; compact customer_pack title, no cover page."""
    sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.492)
    styles=doc.styles
    tokens={'Normal':(11,0,6,'000000',False), 'Title':(24,0,8,'0B2545',True),
            'Subtitle':(10,0,8,'555555',False), 'Heading 1':(16,18,10,'2E74B5',True),
            'Heading 2':(13,14,7,'2E74B5',True), 'Heading 3':(12,10,5,'1F4D78',True),
            'List Bullet':(11,0,4,'000000',False), 'List Number':(11,0,4,'000000',False),
            'Prompt':(10.5,0,6,'000000',False), 'Record Label':(11,6,4,'1F4D78',True),
            'Header':(9,0,0,'555555',False), 'Footer':(9,0,0,'555555',False)}
    for name,(size,before,after,color,bold) in tokens.items():
        st=styles[name] if name in styles else styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(color);st.font.bold=bold
        if name in ('Title','Subtitle','Heading 1','Heading 2','Heading 3','Record Label','Header'):
            st.font.color.rgb=RGBColor(0,0,0)
        pp=st._element.find(qn('w:pPr'))
        if pp is not None:
            for tag in ('w:pBdr','w:shd'):
                for node in list(pp.findall(qn(tag))): pp.remove(node)
        pf=st.paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after);pf.line_spacing=1.25
        pf.widow_control=True
        if name.startswith('Heading') or name in ('Title','Subtitle','Record Label'): pf.keep_with_next=True
    styles['Prompt'].paragraph_format.keep_together=True
    if checklist:
        sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.6)
        sec.left_margin=sec.right_margin=Inches(.7)
        sec.header_distance=sec.footer_distance=Inches(.25)
        for name in ('Normal','List Bullet','List Number','Prompt'):
            styles[name].font.size=Pt(11.5)
            styles[name].paragraph_format.line_spacing=1.12
            styles[name].paragraph_format.space_after=Pt(5)
        styles['Title'].font.size=Pt(25)
        styles['Title'].paragraph_format.space_after=Pt(4)
        styles['Subtitle'].font.size=Pt(10.5)
        styles['Subtitle'].paragraph_format.space_after=Pt(8)
        for name in ('Heading 1','Heading 2','Heading 3'):
            styles[name].font.size=Pt(13)
            styles[name].paragraph_format.space_before=Pt(9)
            styles[name].paragraph_format.space_after=Pt(4)
        st=styles.add_style('Checklist',WD_STYLE_TYPE.PARAGRAPH)
        st.base_style=styles['Normal']
        st.paragraph_format.left_indent=Inches(.29)
        st.paragraph_format.first_line_indent=Inches(-.29)
        st.paragraph_format.keep_together=True
        st.paragraph_format.space_after=Pt(6)
    numbering=doc.part.numbering_part.element
    for abstract_id,fmt,marker in [(70,'bullet','•'),(71,'decimal','%1.')]:
        a=el('abstractNum',abstractNumId=abstract_id);lvl=el('lvl',ilvl=0)
        lvl.extend([el('start',val=1),el('numFmt',val=fmt),el('lvlText',val=marker),el('lvlJc',val='left')])
        pp=el('pPr');tabs=el('tabs');tabs.append(el('tab',val='num',pos=540));pp.append(tabs)
        pp.append(el('ind',left=540,hanging=271));lvl.append(pp);a.append(lvl);numbering.append(a)
    b=el('num',numId=70);b.append(el('abstractNumId',val=70));numbering.append(b)
    sec.header.paragraphs[0].text='SCHEDULING ASSISTANT  |  Package '+VERSION
    foot=sec.footer.paragraphs[0];foot.alignment=2
    foot.add_run('Page ');field=el('fldSimple',instr='PAGE');foot._p.append(field)
    props=doc.core_properties;props.author='';props.last_modified_by='';props.title='Scheduling Assistant'
    props.created=props.modified=METADATA_TIMESTAMP
    props.revision=1


def clean(text):
    def link(m):
        label,target=m.group(1),m.group(2)
        if target.startswith(('http:','https:')): return label+' ('+target+')'
        return LINK_NAMES.get(Path(target.split('#')[0]).name,label)
    text=re.sub(r'\[([^]]+)\]\(([^)]+)\)',link,text)
    text=text.replace('`','')
    return text


def inline(p,text):
    for bit in re.split(r'(\*\*.*?\*\*)',clean(text)):
        if bit.startswith('**') and bit.endswith('**'): p.add_run(bit[2:-2]).bold=True
        else: p.add_run(bit)


def listpara(doc,text,num_id):
    p=doc.add_paragraph(style='List Bullet' if num_id==70 else 'List Number')
    pp=p._p.get_or_add_pPr();num=el('numPr');num.extend([el('ilvl',val=0),el('numId',val=num_id)]);pp.append(num)
    inline(p,text)


def title_text(text):
    text=re.sub(r'\bv\d+(?:\.\d+)+\b', '', clean(text))
    text=re.sub(r'[^\w\s.]',' ',text)
    text=re.sub(r'(?<!\d)\.|\.(?!\d)', ' ', text)
    return re.sub(r'\s+',' ',text).strip()


def write_doc(source, destination, subtitle=None, checklist=False):
    doc=Document();configure(doc,checklist);lines=source.splitlines();i=0;seq=100;checkbox_number=0
    if destination.name in PROFILE_STYLE_DOCS:
        # Long rules text: slightly tighter spacing keeps the document compact.
        doc.styles['Normal'].paragraph_format.space_after=Pt(4)
        doc.styles['List Bullet'].paragraph_format.space_after=Pt(2)
        for style in ('Normal', 'List Bullet', 'List Number'):
            doc.styles[style].paragraph_format.line_spacing=1.2
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1;continue
        if line.startswith('```'):
            i+=1;buf=[]
            while i<len(lines) and not lines[i].startswith('```'): buf.append(lines[i]);i+=1
            prompt=re.sub(r'(?<![.:])\n(?=\S)',' ','\n'.join(buf))
            p=doc.add_paragraph(style='Prompt');p.add_run(prompt);i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=re.split(r'(?<!\\)\|',lines[i].strip().strip('|'))
                if not all(re.fullmatch(r'\s*:?-+:?\s*',v) for v in row): rows.append([v.strip().replace('\\|','|') for v in row])
                i+=1
            headers=rows[0]
            if len(headers)>2 and len(rows)>4 and all(all(re.fullmatch(r'\[.*\]',v) for v in r[1:]) for r in rows[1:]):
                inline(doc.add_paragraph(),'Cover these topics; reference supplied products instead of retyping them:')
                for row in rows[1:]: listpara(doc,row[0],70)
                inline(doc.add_paragraph(style='Record Label'),'Copy this entry for each topic that needs a record')
                for h,v in zip(headers,['[topic]']+rows[1][1:]):
                    inline(doc.add_paragraph(),'**'+h+':** '+v)
                continue
            for row in rows[1:]:
                if len(row)!=len(headers): raise ValueError(('ragged table',destination,row))
                if len(headers)==2:
                    p=doc.add_paragraph();inline(p,'**'+row[0]+':** '+row[1])
                else:
                    p=doc.add_paragraph(style='Record Label');inline(p,headers[0]+': '+row[0])
                    for h,v in zip(headers[1:],row[1:]):
                        p=doc.add_paragraph();inline(p,'**'+h+':** '+v)
            continue
        heading=re.match(r'^(#{1,6}) (.*)',line)
        if heading:
            level=len(heading[1]);title=title_text(heading[2])
            if level==1:
                inline(doc.add_paragraph(style='Title'),title)
                inline(doc.add_paragraph(style='Subtitle'),subtitle or 'Package '+VERSION+' · Human-led scheduling')
            else: inline(doc.add_paragraph(style='Heading '+str(min(level-1,3))),title)
            i+=1;continue
        if re.match(r'^\d+\. ',line):
            seq+=1;n=el('num',numId=seq);n.append(el('abstractNumId',val=71))
            # Each separate numbered list restarts at 1.
            override=el('lvlOverride',ilvl=0);override.append(el('startOverride',val=1));n.append(override)
            doc.part.numbering_part.element.append(n)
            while i<len(lines) and re.match(r'^\d+\. ',lines[i].strip()):
                listpara(doc,re.sub(r'^\d+\. ','',lines[i].strip()),seq);i+=1
            continue
        if line.startswith('- [ ] '):
            checkbox_number+=1
            p=doc.add_paragraph(style='Checklist' if checklist else 'Normal')
            box=p.add_run('☐  ');box.font.name='DejaVu Sans';box.font.size=Pt(12)
            p.add_run(str(checkbox_number)+'. ').bold=True
            inline(p,line[6:]);i+=1;continue
        if line.startswith('- '): listpara(doc,line[2:],70);i+=1;continue
        buf=[line];i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^(#|\||```|- |\d+\. )',lines[i].strip()):
            buf.append(lines[i].strip());i+=1
        inline(doc.add_paragraph(),' '.join(buf))
    destination.parent.mkdir(parents=True,exist_ok=True)
    doc.save(destination)
    with zipfile.ZipFile(destination) as z: contents={n:z.read(n) for n in z.namelist() if not n.startswith('docProps/thumbnail')}
    with zipfile.ZipFile(destination) as z:
        for n in z.namelist():
            if n.startswith('docProps/thumbnail'): contents[n]=z.read(n)
    with zipfile.ZipFile(destination,'w',compression=zipfile.ZIP_STORED) as z:
        for n in sorted(contents):
            info=zipfile.ZipInfo(n,ZIP_TIMESTAMP);info.compress_type=zipfile.ZIP_STORED;z.writestr(info,contents[n])


def make_zip(folder,destination):
    destination.parent.mkdir(exist_ok=True,parents=True)
    with zipfile.ZipFile(destination,'w',compression=zipfile.ZIP_STORED) as z:
        for p in sorted(folder.rglob('*')):
            name=p.relative_to(folder).as_posix()+('/' if p.is_dir() else '')
            info=zipfile.ZipInfo(name,ZIP_TIMESTAMP);info.compress_type=zipfile.ZIP_STORED
            z.writestr(info,b'' if p.is_dir() else p.read_bytes())


def write_text(destination, text):
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(text.encode('utf-8'))


def build(out):
    validate_reproducibility_contract()
    staging=out/'build';downloads=out/'downloads'
    if staging.exists(): shutil.rmtree(staging)
    if downloads.exists(): shutil.rmtree(downloads)
    replaceable=staging/'replaceable'
    write_doc(START_SOURCE.read_text(),replaceable/START_DOC,checklist=True)
    for source,title in TASKS:
        text=task_file_text(source,title);validate_task_file(text,title)
        write_text(replaceable/TASK_DIR/(title+'.txt'),text)

    full=staging/'full'/KIT
    shutil.copytree(replaceable,full)
    write_doc(squadron_rules_source(),full/RULES_DOC,'Approved local rules - the DO approves changes')
    week=full/WEEK_TEMPLATE
    for directory in WEEKLY_DIRS:
        (week/directory).mkdir(parents=True)
    write_doc(INPUT_SOURCE.read_text(),week/INPUT_CHECKLIST,checklist=True)
    make_zip(staging/'full',downloads/FULL_ZIP)

    update=staging/'update'
    shutil.copytree(replaceable,update)
    if actionable_migrations():
        write_doc(persistent_migration_source(),update/MIGRATION_DOC,'Read, act on what applies, then delete',checklist=True)
    make_zip(update,downloads/UPDATE_ZIP)

    manifests={}
    for name in ZIPS:
        p=downloads/name
        with zipfile.ZipFile(p) as z: entries=z.namelist()
        manifests[name]={'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'files':entries}
    (downloads/'manifest.json').write_text(json.dumps({'package_version':VERSION,'packages':manifests},indent=2)+'\n')
    print(json.dumps({p.name:p.stat().st_size for p in sorted(downloads.iterdir())},indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,required=True)
    build(parser.parse_args().output.resolve())
