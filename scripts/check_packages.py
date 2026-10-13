"""Static package checks. These do not test GenAI.mil model behavior."""
from pathlib import Path, PurePosixPath
from io import BytesIO
import hashlib
import json
import re
import shutil
import sys
import tempfile
import zipfile
from docx import Document

from build_packages import (
    ROOT, VERSION, KIT, FULL_ZIP, UPDATE_ZIP, ZIPS, TASK_DIR, TASKS, WEEK_TEMPLATE, WEEKLY_DIRS,
    START_DOC, RULES_DOC, INPUT_CHECKLIST, MIGRATION_DOC, CORE_SOURCE, WEEK_FILE_SOURCE, START_SOURCE,
    INPUT_SOURCE, METADATA_TIMESTAMP, ZIP_TIMESTAMP, NAMED_MODEL, actionable_migrations, clean, title_text,
    persistent_migration_source, squadron_rules_source, task_file_text, validate_task_file,
    validate_reproducibility_contract,
)
from check_persistent_migrations import validate_current as validate_persistent

# Internal codes and terms that must never appear in what operators read directly.
OPERATOR_JARGON = re.compile(
    r'P-#|\b[PFRDWXSC]-\d{3}\b|\bledger\b|\bbaseline\b|\bprimer\b|CANNOT VERIFY|SOURCE-BACKED|'
    r'INTEGRITY FAILED|PRODUCT GATHERING|subagent', re.I)
COMPACT_DOCS = (START_DOC, INPUT_CHECKLIST, MIGRATION_DOC)


def text(docx_bytes):
    doc=Document(BytesIO(docx_bytes))
    return '\n'.join(p.text for p in doc.paragraphs)


def task_names():
    return [title + '.txt' for _, title in TASKS]


def replaceable_inventory():
    return {START_DOC} | {TASK_DIR + '/' + name for name in task_names()}


def full_inventory():
    files = {KIT + '/' + path for path in replaceable_inventory()}
    files |= {KIT + '/' + RULES_DOC, KIT + '/' + WEEK_TEMPLATE + '/' + INPUT_CHECKLIST}
    return files


def update_inventory():
    return replaceable_inventory() | ({MIGRATION_DOC} if actionable_migrations() else set())


def validate_full_inventory(paths):
    expected = full_inventory()
    if set(paths) != expected:
        raise ValueError('Full kit must contain exactly the operator files: ' + str(sorted(set(paths) ^ expected)))


def validate_update_inventory(paths):
    for path in paths:
        if path == RULES_DOC or path.startswith(('WEEKS/', 'Squadron Rules')):
            raise ValueError('Update must never contain Squadron Rules or weekly work: ' + path)
    expected = update_inventory()
    if set(paths) != expected:
        raise ValueError('Update must contain only replaceable files: ' + str(sorted(set(paths) ^ expected)))


def jargon_free(name, rendered):
    match = OPERATOR_JARGON.search(rendered)
    if match:
        raise ValueError(f'Operator document {name} uses internal jargon: {match.group(0)!r}')


def normalize(s):
    s=re.sub(r'☐\s*\d+\.\s*','',s)
    return re.sub(r'\s+',' ',clean(s).replace('**','').replace('[ ] ','')).strip()


def _table_row(line):
    return [v.strip().replace('\\|','|') for v in re.split(r'(?<!\\)\|',line.strip().strip('|'))]


def _table_blocks(source):
    lines=source.splitlines();i=0
    while i<len(lines):
        if not lines[i].strip().startswith('|'):
            i+=1;continue
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            block.append(lines[i]);i+=1
        rows=[]
        for line in block:
            cells=_table_row(line)
            if all(re.fullmatch(r'\s*:?-+:?\s*',v) for v in cells):
                continue
            rows.append(cells)
        if rows:
            yield rows


def source_coverage(source, actual):
    """Require maintained prose, prompts, headers and meaningful table content."""
    actual=normalize(actual)
    for raw in source.splitlines():
        s=raw.strip()
        if not s or s.startswith(('```','|')): continue
        if re.match(r'^#{1,6} ',s): s=title_text(re.sub(r'^#{1,6} ','',s))
        s=re.sub(r'^(#{1,6} |[-] |\d+\. )','',s)
        fragment=normalize(s)
        if fragment:
            assert fragment in actual, ('Missing source text',s[:140])
    for prompt in re.findall(r'```[^\n]*\n(.*?)\n```',source,re.S):
        fragment=normalize(prompt)
        if fragment:
            assert fragment in actual, ('Missing prompt content',fragment[:140])
    for rows in _table_blocks(source):
        width=len(rows[0])
        for row_index,row in enumerate(rows):
            assert len(row)==width, ('Ragged source table',row)
            if width==2 and row_index==0:
                continue
            for cell in row:
                if re.fullmatch(r'\[[^]]+\]',cell.strip()):
                    continue
                fragment=normalize(cell)
                if fragment:
                    assert fragment in actual, ('Missing source table cell',cell[:140])


def _datetime_tuple(value):
    return (value.year,value.month,value.day,value.hour,value.minute,value.second)


def check(out):
    validate_persistent()
    validate_reproducibility_contract()
    assert ZIP_TIMESTAMP==(1980,1,1,0,0,0)
    assert _datetime_tuple(METADATA_TIMESTAMP)==ZIP_TIMESTAMP
    manifest=json.loads((out/'downloads/manifest.json').read_text())
    assert manifest['package_version']==VERSION
    assert set(manifest['packages'])==set(ZIPS)
    contents={}
    for name,entry in manifest['packages'].items():
        p=out/'downloads'/name
        assert hashlib.sha256(p.read_bytes()).hexdigest()==entry['sha256']
        with zipfile.ZipFile(p) as z:
            assert z.testzip() is None
            assert z.namelist()==entry['files']
            assert len(set(z.namelist()))==len(z.namelist())
            assert all(info.compress_type==zipfile.ZIP_STORED for info in z.infolist()), ('Outer ZIP compression drift',name)
            assert all(info.date_time==ZIP_TIMESTAMP for info in z.infolist()), ('Outer ZIP timestamp drift',name)
            contents[name]={n:z.read(n) for n in z.namelist() if not n.endswith('/')}
            directories={n for n in z.namelist() if n.endswith('/')}
        if name==FULL_ZIP:
            for d in WEEKLY_DIRS:
                assert KIT+'/'+WEEK_TEMPLATE+'/'+d+'/' in directories, ('Missing weekly folder',d)
        for n,data in contents[name].items():
            assert '..' not in PurePosixPath(n).parts and not n.startswith('/')
            assert n.endswith(('.docx','.txt')),n
            if n.endswith('.txt'):
                validate_task_file(data.decode('utf-8'),n)
                continue
            with zipfile.ZipFile(BytesIO(data)) as z:
                assert z.testzip() is None
                assert not any('vbaProject' in x for x in z.namelist())
                assert all(info.compress_type==zipfile.ZIP_STORED for info in z.infolist()), ('DOCX compression drift',n)
                assert all(info.date_time==ZIP_TIMESTAMP for info in z.infolist()), ('DOCX timestamp drift',n)
            d=Document(BytesIO(data));sec=d.sections[0]
            assert sec.page_width.twips==12240 and sec.page_height.twips==15840
            base=PurePosixPath(n).name
            compact=base in COMPACT_DOCS
            margins=tuple(x.twips for x in (sec.left_margin,sec.right_margin,sec.top_margin,sec.bottom_margin))
            assert margins==((1008,1008,936,864) if compact else (1440,1440,1440,1440)), (n,margins)
            assert d.styles['Normal'].font.size.pt==(11.5 if compact else 11)
            spacing=1.12 if compact else (1.2 if base==RULES_DOC else 1.25)
            assert abs(d.styles['Normal'].paragraph_format.line_spacing-spacing)<.002
            for style in ('Title','Subtitle','Heading 1','Heading 2','Heading 3','Header'):
                assert str(d.styles[style].font.color.rgb)=='000000', (n,style)
            assert not d.tables
            assert d.core_properties.created and _datetime_tuple(d.core_properties.created)==ZIP_TIMESTAMP, ('DOCX created metadata drift',n)
            assert d.core_properties.modified and _datetime_tuple(d.core_properties.modified)==ZIP_TIMESTAMP, ('DOCX modified metadata drift',n)
            rendered=text(data)
            assert '.md' not in rendered, ('Operator Markdown reference',n)
            assert not NAMED_MODEL.search(rendered), ('Named model in current package',n)

    full=contents[FULL_ZIP];update=contents[UPDATE_ZIP];k=KIT+'/'
    validate_full_inventory(full)
    validate_update_inventory(update)

    core=CORE_SOURCE.read_text().strip();week=WEEK_FILE_SOURCE.read_text().strip()
    for source,title in TASKS:
        path=TASK_DIR+'/'+title+'.txt'
        expected=task_file_text(source,title).encode('utf-8')
        assert full[k+path]==update[path]==expected, ('Task file differs from its sources',title)
        rendered=expected.decode('utf-8')
        assert core in rendered and week in rendered, ('Task file lost the shared core or WEEK FILE format',title)
        assert ('THIS TASK: '+title) in rendered
    assert full[k+START_DOC]==update[START_DOC]

    source_coverage(START_SOURCE.read_text(),text(update[START_DOC]))
    source_coverage(INPUT_SOURCE.read_text(),text(full[k+WEEK_TEMPLATE+'/'+INPUT_CHECKLIST]))
    source_coverage(squadron_rules_source(),text(full[k+RULES_DOC]))
    jargon_free(START_DOC,text(update[START_DOC]))
    jargon_free(INPUT_CHECKLIST,text(full[k+WEEK_TEMPLATE+'/'+INPUT_CHECKLIST]))
    assert sum(p.text.startswith('☐') for p in Document(BytesIO(full[k+WEEK_TEMPLATE+'/'+INPUT_CHECKLIST])).paragraphs)==INPUT_SOURCE.read_text().count('- [ ] ')
    if actionable_migrations():
        source_coverage(persistent_migration_source(),text(update[MIGRATION_DOC]))

    # Simulate START HERE's update and a rollback. Squadron Rules and weekly work must survive both.
    with tempfile.TemporaryDirectory() as temp:
        base=Path(temp)
        with zipfile.ZipFile(out/'downloads'/FULL_ZIP) as z:z.extractall(base)
        install=base/KIT
        prior={n:data for n,data in update.items() if n==START_DOC or n.startswith(TASK_DIR+'/')}
        samples={RULES_DOC:b'LOCAL APPROVED RULES',
                 'WEEKS/Week of 2099-01-04/WEEK FILE 03 - Thu QC.txt':b'HUMAN DECISIONS',
                 'WEEKS/Week of 2099-01-04/Schedules/Published.xlsx':b'WEEKLY BASELINE',
                 WEEK_TEMPLATE+'/Inputs/custom-owner-notes.txt':b'KEEP NOTES'}
        for name,data in samples.items():
            p=install/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(data)
        (install/TASK_DIR/'0 OLD - retired task.txt').write_bytes(b'old')
        # Update: delete TASK FILES, then copy everything in Update.zip into the kit folder.
        shutil.rmtree(install/TASK_DIR)
        for name,data in update.items():
            p=install/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(data)
        assert not (install/TASK_DIR/'0 OLD - retired task.txt').exists()
        for name,data in samples.items():assert (install/name).read_bytes()==data
        for name in update:assert (install/name).read_bytes()==update[name]
        # Rollback: delete TASK FILES and restore a prior release's replaceable files.
        (install/TASK_DIR/'9 NEWER - broken.txt').write_bytes(b'broken')
        shutil.rmtree(install/TASK_DIR)
        for name,data in prior.items():
            p=install/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_bytes(data)
        assert not (install/TASK_DIR/'9 NEWER - broken.txt').exists()
        for name,data in samples.items():assert (install/name).read_bytes()==data
    print('PASS: ZIP integrity/stored format; deterministic sentinel metadata; task-file contract; shared core in every task; source coverage; operator jargon; update-only-replaceable; update/rollback preservation.')


if __name__=='__main__':check(Path(sys.argv[1]))
